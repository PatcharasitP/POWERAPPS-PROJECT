"""สร้างไฟล์เทมเพลต LOG ELECTRIC (.xlsx) จาก web/config.js แล้วฝังเป็น base64 ใน web/template.js (web resource เก็บ .xlsx ตรงไม่ได้)
เทมเพลตมีเฉพาะช่องที่ต้องกรอก (ตัด AUTO) dropdown ผูกกับชีต Lists วันที่บังคับเป็นวันที่ AMOUNT บังคับเป็นตัวเลข
รัน: python3 tools/make_template.py"""
import base64, datetime, io, json, subprocess, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
cfg = json.loads(subprocess.check_output(["node", "-e",
    "import('./web/config.js').then(m=>console.log(JSON.stringify({f:m.INPUT_FIELDS,sheet:m.TEMPLATE_SHEET,co:m.COMPANIES,ex:m.TEMPLATE_EXAMPLE})))"], cwd=ROOT, text=True))
fields, SHEET = cfg["f"], cfg["sheet"]
EXAMPLE = cfg["ex"]  # ตัวอย่างครบทุกคอลัมน์ ตามแบบที่พี่ปอนด์ส่งมา 29/10 รอบ 4
ROWS = 5000

wb = Workbook()
ws = wb.active
ws.title = SHEET
lists = wb.create_sheet("Lists")
lists.sheet_state = "hidden"
guide = wb.create_sheet("How to fill")

red = PatternFill("solid", fgColor="FDECEC")
grey = PatternFill("solid", fgColor="F2F2F2")
lc = 0
for i, f in enumerate(fields, 1):
    col = get_column_letter(i)
    c = ws.cell(1, i, f["t"])
    c.font = Font(bold=True, color="A80000" if f["input"] == "required" else "5F5B57")
    c.fill = red if f["input"] == "required" else grey
    c.alignment = Alignment(vertical="center")
    ws.column_dimensions[col].width = max(14, len(f["t"]) + 4)
    rng = f"{col}2:{col}{ROWS + 1}"
    if f["kind"] == "choice":
        lc += 1
        L = get_column_letter(lc)
        lists.cell(1, lc, f["t"])
        for j, o in enumerate(f["options"], 2):
            lists.cell(j, lc, o).number_format = "@"
        dv = DataValidation(type="list", formula1=f"=Lists!${L}$2:${L}${len(f['options']) + 1}", allow_blank=f["input"] != "required",
                            showErrorMessage=True, errorTitle=f["t"], error="Choose a value from the list.")
        if f["key"] == "companyCode":
            ws.column_dimensions[col].number_format = "@"
    elif f["kind"] == "date":
        dv = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)", allow_blank=True,
                            showErrorMessage=True, errorTitle=f["t"], error="Enter a date, e.g. 29/10/2026.")
        ws.column_dimensions[col].number_format = "dd/mm/yyyy"
    elif f["kind"] == "num":
        dv = DataValidation(type="decimal", operator="between", formula1="-99999999999", formula2="99999999999", allow_blank=True,
                            showErrorMessage=True, errorTitle=f["t"], error="Enter a number.")
        ws.column_dimensions[col].number_format = "#,##0.00"
    else:
        dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(f.get("max", 255)), allow_blank=True,
                            showErrorMessage=True, errorTitle=f["t"], error=f"Up to {f.get('max', 255)} characters.")
        ws.column_dimensions[col].number_format = "@"
    dv.add(rng)
    ws.add_data_validation(dv)
ws.freeze_panes = "A2"


def typed(f, v):
    """ค่าตัวอย่างเป็นชนิดจริงของช่อง วันที่เป็นวันที่ เงินเป็นตัวเลข ที่เหลือเป็นข้อความ (COMPANY CODE ด้วย)"""
    if f["kind"] == "date":
        d, m, y = map(int, v.split("/"))
        return datetime.datetime(y, m, d), "dd/mm/yyyy"
    if f["kind"] == "num":
        return float(v), "0.00"
    return v, "@"


# แถวตัวอย่าง 1 แถวในชีตข้อมูล (พี่ปอนด์ 29/10 รอบ 4) ตอนนำเข้า แถวที่ค่าตรงตัวอย่างทุกช่องถูกข้ามให้เอง
for i, f in enumerate(fields, 1):
    v, fmt = typed(f, EXAMPLE[f["key"]])
    c = ws.cell(2, i, v)
    c.number_format = fmt

guide.append(["Column", "Required", "Type", "Options", "Example"])
for c in guide[1]:
    c.font = Font(bold=True)
TYPE = {"date": "Date (dd/mm/yyyy)", "num": "Number", "choice": "Text (Dropdown)", "text": "Text"}
missing = [f["key"] for f in fields if f["key"] not in EXAMPLE]
assert not missing, f"ไม่มีตัวอย่าง {missing}"
for r, f in enumerate(fields, 2):
    guide.append([f["t"], "Yes" if f["input"] == "required" else "No", TYPE[f["kind"]],
                  ", ".join(f["options"]) if f["kind"] == "choice" else "", None])
    for c in guide[r][:4]:
        c.number_format = "@"
    v, fmt = typed(f, EXAMPLE[f["key"]])
    guide.cell(r, 5, v).number_format = fmt
for col, w in zip("ABCDE", (26, 10, 20, 28, 26)):
    guide.column_dimensions[col].width = w

buf = io.BytesIO()
wb.save(buf)
data = buf.getvalue()
(ROOT / "web" / "template.js").write_text(
    "// สร้างโดย tools/make_template.py จาก config.js ห้ามแก้มือ\n"
    f"export const TEMPLATE_NAME = \"LOG_ELECTRIC_TEMPLATE.xlsx\";\nexport const TEMPLATE_B64 = \"{base64.b64encode(data).decode()}\";\n", encoding="utf-8")
(ROOT / "tools" / "out").mkdir(exist_ok=True)
(ROOT / "tools" / "out" / "LOG_ELECTRIC_TEMPLATE.xlsx").write_bytes(data)
print("template", len(data), "bytes, columns", [f["t"] for f in fields])
