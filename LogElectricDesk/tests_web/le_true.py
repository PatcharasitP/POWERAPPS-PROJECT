#!/usr/bin/env python3
"""le_true.py  ใช้งานแบบที่ทรูจะใช้จริง (พี่ปอนด์ 30/10/2026 "จัดเทสพวกการใช้งานสำหรับที่ทรูด้วยนะ")

T1 วันทำงานของลูกทีม (64 เป็น LE Member): โหลดเทมเพลตจากแอป กรอกแบบคนไทย (ชื่อไทย , ปี พ.ศ. , เงินมีลูกน้ำ) นำเข้า แก้ 1 แถว ส่งออกแถวของตัวเอง
T2 หัวหน้า (65 เป็น LE Admin): เลือกชื่อลูกทีมในช่อง All people เห็นเท่าใน Dataverse จัดกลุ่มตาม TEAM ส่งออกเฉพาะคนนั้น
T3 นำเข้ายอดทั้งเดือน 500 แถวในไฟล์เดียว ลงครบ ไม่ซ้ำ INDEX
T4 จอโน้ตบุ๊กบริษัทในกรอบ MDA (1366 , 1280 ซูม 150% , 1536 ซูม 125% , 1920) แถบเครื่องมือ ตาราง ฟอร์ม กล่องตัวกรอง หน้าต่างส่งออก
T5 Firefox (เผื่อเครื่องที่ไม่ใช้ Edge หรือ Chrome) เปิด ค้น กรอง ส่งออก ตรวจไฟล์นำเข้า
แถวที่เทสสร้าง PRO NUMBER ขึ้นต้น BUG-T ลบทิ้งตอนจบ (cleanup ของ le_bughunt)
ใช้: PVDESK_TOKEN=$(az account get-access-token --resource https://pvdesk.crm5.dynamics.com --query accessToken -o tsv) .venv/bin/python Log_Electric/desk/tests_web/le_true.py [T1 T2 T3 T4 T5]
"""
import datetime, pathlib, sys, time, urllib.parse
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from le_ui import check, RESULTS, EVID, FILES, dv, count, user_id, Server, open_page  # noqa: E402
from le_bughunt import TMP, HEADS, row, xlsx, upload, dl, toast, save, open_row, hist_n, cleanup, pick_owner  # noqa: E402

GROUPS = set(sys.argv[1:]) or {"T1", "T2", "T3", "T4", "T5"}
THAI_NAMES = ["บริษัท ทรัพย์สมบูรณ์ จำกัด", "นางสาวสมหญิง ใจดี", "หจก. ไฟฟ้าเจริญ", "นายสมชาย รักไทย", "การไฟฟ้าส่วนภูมิภาค สาขาบางพลี"]


def bug_count(flt=""):
    return count(flt="startswith(le_pronumber,'BUG-T')" + (f" and {flt}" if flt else ""))


def import_click(p, timeout=300000):
    p.click("#btn-confirm-import"); p.click("dialog.dlg[open] .btn.dark")  # แอปถามยืนยันทุกครั้ง (Confirm import)
    p.wait_for_selector("#imp-done, #imp-failed", timeout=timeout)
    return p.inner_text("#imp-result")


def t1(b, srv, u64):
    ctx, p = open_page(b, srv, caller=u64)
    before = count(flt=f"_ownerid_value eq {u64}")
    p.click("#tab-import")
    with p.expect_download() as d:
        p.click("#btn-template")
    tpl = dl(d)
    wb = load_workbook(tpl); ws = wb["LOG ELECTRIC"]
    # กรอกต่อจากแถวตัวอย่าง แบบที่ลูกทีมทำ: ปี พ.ศ. เป็นข้อความ , เงินมีลูกน้ำ , ชื่อไทย , ช่องหมายเหตุไทย
    heads = [c.value for c in ws[1]]
    for i in range(25):
        v = {"LOG BOOK DATE": f"{(i % 28) + 1:02d}/10/2569", "PRO NUMBER": f"BUG-T1-{i:02d}", "SITE CODE CANDIDATE": "BKK4874A",
             "NAME": THAI_NAMES[i % 5], "DUE DATE": datetime.datetime(2026, 11, (i % 28) + 1), "AMOUNT": f"{1250 + i:,}.50",
             "PAYMENT METHOD": "BANK TRANSFER" if i % 2 else "CHEQUE", "ACCOUNTANT": "สมศรี บัญชี", "REQUESTOR": "ปอนด์ ทดสอบ",
             "TEAM": ["IBC", "MACRO", "ONLINE"][i % 3], "SYSTEM": "SITE RENTAL", "COMPANY CODE": "1068", "REMARK": f"มิเตอร์ {i + 1} หน่วย"}
        ws.append([v.get(h) for h in heads])
    filled = TMP / "true_t1.xlsx"; wb.save(filled)
    upload(p, filled)
    ok = p.is_visible("#imp-ready")
    check("T1 ลูกทีมกรอกเทมเพลตจากแอป (ปี พ.ศ. , เงินมีลูกน้ำ , ชื่อไทย) ไฟล์ผ่าน แถวตัวอย่างถูกข้าม", ok and "25 rows" in p.inner_text("#imp-result"),
          p.inner_text("#imp-result")[:160].replace("\n", " "))
    res = import_click(p)
    after = count(flt=f"_ownerid_value eq {u64}")
    got = dv("GET", "le_logelectrics?$select=le_logbookdate,le_amount,le_name,le_index,_ownerid_value&$orderby=le_pronumber&$filter="
             + urllib.parse.quote("startswith(le_pronumber,'BUG-T1-')"))["value"]
    idx = [r["le_index"] for r in got]
    check("T1 นำเข้าแล้วลงครบ 25 แถว เป็นของลูกทีมเอง", after - before == 25 and all(r["_ownerid_value"] == u64 for r in got), f"{before} -> {after} / {res[:60]}")
    check("T1 ปี 2569 กลายเป็น 2026 เงิน 1,250.50 เป็น 1250.5 ชื่อไทยครบ", got and got[0]["le_logbookdate"] == "2026-10-01" and float(got[0]["le_amount"]) == 1250.5
          and got[0]["le_name"] == THAI_NAMES[0], str({k: got[0].get(k) for k in ("le_logbookdate", "le_amount", "le_name")}) if got else "ไม่มีแถว")
    check("T1 INDEX ไม่ซ้ำ", len(set(idx)) == len(idx) == 25, str(idx[:5]))
    p.click("#tab-records"); p.wait_for_timeout(400)
    summ = p.inner_text("#rec-summary")
    check("T1 หน้าตารางของลูกทีมเห็นแถวตัวเองครบทันที (ไม่ต้องกด Refresh)", summ.startswith(f"{after} of"), summ)
    # แก้ 1 แถว ประวัติขึ้น 1 รายการ
    open_row(p, "BUG-T1-03"); p.fill("#f-amount", "9999"); msg = save(p)
    rid = dv("GET", "le_logelectrics?$select=le_logelectricid,le_amount&$filter=le_pronumber eq 'BUG-T1-03'")["value"][0]
    check("T1 ลูกทีมแก้เงินแถวตัวเองได้ ประวัติขึ้น 1 รายการ", float(rid["le_amount"]) == 9999 and hist_n(rid["le_logelectricid"]) >= 1, f"{msg} / {rid['le_amount']}")
    # ส่งออก: ได้เฉพาะแถวของตัวเอง ตัวไทยไม่เพี้ยน
    p.fill("#rec-search", ""); p.wait_for_timeout(400)
    p.click("#btn-export"); p.wait_for_selector("dialog.dlg[open] #export-go")
    with p.expect_download() as d:
        p.click("#export-go")
    f = dl(d); ws2 = load_workbook(f).active
    rows = [r for r in ws2.iter_rows(min_row=2, values_only=True) if any(r)]
    hd = [c.value for c in ws2[1]]
    names = {r[hd.index("NAME")] for r in rows}
    check("T1 ลูกทีมส่งออกได้เฉพาะแถวของตัวเอง ตัวไทยอ่านออก", len(rows) == after and THAI_NAMES[2] in names, f"{len(rows)} แถว ของตัวเอง {after}")
    p.screenshot(path=str(EVID / "true_t1_member.png"))
    check("T1 ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
    ctx.close()


def t2(b, srv, u64):
    ctx, p = open_page(b, srv)
    n64 = count(flt=f"_ownerid_value eq {u64}")
    pick_owner(p, u64)
    summ = p.inner_text("#rec-summary")
    check("T2 หัวหน้าเลือกชื่อลูกทีมในช่อง All people เห็นแถวเท่าใน Dataverse", summ.startswith(f"{n64} of "), f"{summ} / Dataverse {n64}")  # แบบเดียวกับตัวกรอง n of ทั้งหมด
    p.select_option("#rec-group", "team"); p.wait_for_timeout(400)
    groups = p.locator("#rec-table tr.grp").count()
    p.screenshot(path=str(EVID / "true_t2_manager.png"))
    check("T2 จัดกลุ่มตาม TEAM ระหว่างเลือกคนอยู่ ได้กลุ่มตามทีมของคนนั้น", groups >= 1, f"{groups} กลุ่ม")
    p.select_option("#rec-group", ""); p.wait_for_timeout(300)
    p.click("#btn-export"); p.wait_for_selector("dialog.dlg[open] #export-go")
    with p.expect_download() as d:
        p.click("#export-go")
    ws = load_workbook(dl(d)).active
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]
    check("T2 ส่งออกตอนเลือกคน ได้เฉพาะแถวของคนนั้น", len(rows) == n64, f"{len(rows)} / {n64}")
    pick_owner(p, "")
    check("T2 กลับเป็น All people เห็นทุกแถว", p.inner_text("#rec-summary").startswith(f"{count()} of"), p.inner_text("#rec-summary"))
    check("T2 ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
    ctx.close()


def t3(b, srv):
    ctx, p = open_page(b, srv)
    before = bug_count()
    path = xlsx("true_t3_500.xlsx", [row(i, **{"PRO NUMBER": f"BUG-T3-{i:03d}", "NAME": THAI_NAMES[i % 5]}) for i in range(500)])
    upload(p, path)
    ok = p.is_visible("#imp-ready")
    t0 = time.time()
    res = import_click(p, timeout=600000)
    sec = time.time() - t0
    n = bug_count("startswith(le_pronumber,'BUG-T3-')")
    idx = [r["le_index"] for r in dv("GET", "le_logelectrics?$select=le_index&$filter=" + urllib.parse.quote("startswith(le_pronumber,'BUG-T3-')"))["value"]]
    check("T3 นำเข้า 500 แถวในไฟล์เดียว ลงครบ 500 ภายใน 3 นาที", ok and n == 500, f"{sec:.0f} วินาที ลง {n} / {res[:80]}")
    check("T3 นำเข้า 500 แถว ภายใน 3 นาที", sec <= 180, f"{sec:.0f} วินาที")
    check("T3 INDEX ของ 500 แถวไม่ซ้ำ", len(set(idx)) == len(idx) == 500, f"{len(set(idx))} ค่าไม่ซ้ำ")
    p.click("#tab-records"); p.wait_for_timeout(500)
    check("T3 หน้าตารางนับรวมแถวใหม่ทันที", p.inner_text("#rec-summary").startswith(f"{count()} of"), f"{p.inner_text('#rec-summary')} / {before} + 500")
    check("T3 ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
    ctx.close()


# ขนาดจอในกรอบ MDA: หน้าต่างเบราว์เซอร์ลบแถบแท็บกับแถบที่อยู่ (ราว 110) และหัว MDA (48)
LAPTOPS = [("1366x768 100%", 1366, 610), ("1920x1080 150%", 1280, 560), ("1920x1080 125%", 1536, 700), ("1920x1080 100%", 1920, 920)]
INSIDE = """(sel) => { const e = document.querySelector(sel); if (!e) return 'missing ' + sel; const b = e.getBoundingClientRect();
  return b.left >= -1 && b.right <= innerWidth + 1 && b.top >= -1 && b.bottom <= innerHeight + 1 ? true : sel + ' ' + [b.left, b.top, b.right, b.bottom].map(Math.round); }"""


def t4(b, srv):
    for label, w, hgt in LAPTOPS:
        ctx, p = open_page(b, srv, vp=(w, hgt))
        pre = f"T4 [{label} กรอบ {w}x{hgt}]"
        lines = p.evaluate("""() => new Set([...document.querySelector('.toolbar').children].filter((e) => e.offsetParent && e.getBoundingClientRect().height > 0)
          .map((e) => Math.round(e.getBoundingClientRect().top / 20))).size""")
        check(f"{pre} แถบเครื่องมืออยู่บรรทัดเดียว ไม่ล้นแนวนอน", lines == 1 and p.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"), f"{lines} บรรทัด")
        seen = p.evaluate("""() => { const w = document.querySelector('.table-wrap').getBoundingClientRect();
          return [...document.querySelectorAll('.table-wrap tbody tr[data-id]')].filter((r) => { const b = r.getBoundingClientRect(); return b.top >= w.top && b.bottom <= Math.min(w.bottom, innerHeight); }).length; }""")
        check(f"{pre} เห็นข้อมูลในตารางอย่างน้อย 8 แถวโดยไม่ต้องเลื่อน", seen >= 8, f"{seen} แถว")
        p.screenshot(path=str(EVID / f"true_t4_{w}x{hgt}.png"))
        p.click('.xf-btn[data-col="name"]'); p.wait_for_selector(".xf-pop:not([hidden]) .xf-item"); p.wait_for_timeout(250)
        r = p.evaluate(INSIDE, ".xf-pop")
        check(f"{pre} กล่องตัวกรองอยู่ในจอ", r is True, str(r))
        p.keyboard.press("Escape")
        p.click("#rec-table tbody tr[data-id]"); p.click("#btn-edit"); p.wait_for_selector("#rec-form"); p.wait_for_timeout(300)
        r = p.evaluate(INSIDE, "#btn-save")
        check(f"{pre} ฟอร์มแก้แถว ปุ่ม Save อยู่ในจอ", r is True, str(r))
        p.keyboard.press("Escape"); p.wait_for_timeout(200)
        p.evaluate("document.querySelector('.side-panel.open .sp-close')?.click()"); p.wait_for_timeout(250)
        p.click("#btn-export"); p.wait_for_selector("dialog.dlg[open] #export-go"); p.wait_for_timeout(200)
        r = p.evaluate(INSIDE, "#export-go")
        check(f"{pre} หน้าต่างส่งออก ปุ่มส่งออกเห็นโดยไม่ต้องเลื่อน", r is True, str(r))
        check(f"{pre} ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
        ctx.close()


def t5(pw, srv):
    b = pw.firefox.launch()
    ctx, p = open_page(b, srv)
    n = count()
    check("T5 Firefox เปิดแอปเห็นทุกแถว", p.inner_text("#rec-summary").startswith(f"{n} of {n}"), p.inner_text("#rec-summary"))
    p.fill("#rec-search", "PRO100005"); p.wait_for_timeout(500)
    check("T5 Firefox ค้นหาได้", p.inner_text("#rec-summary").startswith("1 of"), p.inner_text("#rec-summary"))
    p.fill("#rec-search", ""); p.wait_for_timeout(300)
    p.click('.xf-btn[data-col="team"]'); p.wait_for_selector(".xf-pop:not([hidden]) .xf-item")
    p.uncheck(".xf-pop .xf-allrow input"); p.check('.xf-pop .xf-list input[data-v="IBC"]'); p.click(".xf-pop .xf-foot button:last-child")
    p.wait_for_timeout(400)
    ibc = count(flt="le_team eq 'IBC'")
    check("T5 Firefox ตัวกรองแบบ Excel ใช้ได้", p.inner_text("#rec-summary").startswith(f"{ibc} of"), f"{p.inner_text('#rec-summary')} / IBC {ibc}")
    p.click("#btn-export"); p.wait_for_selector("dialog.dlg[open] #export-go")
    with p.expect_download() as d:
        p.click("#export-go")
    ws = load_workbook(dl(d)).active
    check("T5 Firefox ส่งออก XLSX เปิดได้ แถวตรงกับที่กรอง", sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if any(r)) == ibc)
    upload(p, FILES / "bad.xlsx")
    p.wait_for_selector("#imp-rejected", timeout=30000)
    with p.expect_download() as d:
        p.click("#btn-errfile")
    check("T5 Firefox ไฟล์ผิดได้ไฟล์ error กลับ", dl(d).stat().st_size > 1000)
    p.screenshot(path=str(EVID / "true_t5_firefox.png"))
    check("T5 Firefox ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
    ctx.close(); b.close()


def main():
    srv = Server()
    u64 = user_id("64050580")
    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch()
            for g, fn in (("T1", lambda: t1(b, srv, u64)), ("T2", lambda: t2(b, srv, u64)), ("T3", lambda: t3(b, srv)), ("T4", lambda: t4(b, srv)), ("T5", lambda: t5(pw, srv))):
                if g not in GROUPS:
                    continue
                try:
                    fn()
                except Exception as e:
                    msg = str(e); wait = [l.strip() for l in msg.splitlines() if "waiting for" in l][:1]
                    check(f"{g} รันจบไม่สะดุด", False, f"{type(e).__name__}: {msg.splitlines()[0][:150]} {wait}")
            b.close()
    finally:
        srv.proc.terminate()
        cleanup()
    bad = [n for n, ok in RESULTS if not ok]
    print(f"\nรวม {len(RESULTS)} ข้อ ผ่าน {len(RESULTS) - len(bad)} ตก {len(bad)}")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
