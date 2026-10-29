#!/usr/bin/env python3
"""le_bughunt.py  ไล่หาบัคของ Log Electric Desk ด้วยเคสขอบ (พี่ปอนด์ 29/10/2026 "ทำ TEST CASE ทดสอบบัคไปพลาง ๆ")
กลุ่ม A นำเข้าไฟล์แปลก , B ฟอร์ม , C สิทธิ์ (Manager ผ่านหน้าจอจริง) , D ส่งออก , E ความเร็ว 30,000 แถว (ข้อมูลจำลองผ่าน route) , F มือถือ
แถวที่เทสสร้างขึ้นต้น PRO NUMBER ด้วย BUG- และถูกลบตอนจบ บทบาทที่กลุ่ม C ใส่ให้ 64 ถูกถอดคืนตาม baseline ที่จดไว้
ใช้: PVDESK_TOKEN=$(az account get-access-token --resource https://pvdesk.crm5.dynamics.com --query accessToken -o tsv) .venv/bin/python Log_Electric/desk/tests_web/le_bughunt.py [A B C D E F]
"""
import csv, datetime, io, json, os, pathlib, sys, time, urllib.error, urllib.parse, urllib.request
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from le_ui import check, RESULTS, EVID, dv, count, user_id, Server, open_page, BASE  # noqa: E402

TMP = pathlib.Path("/tmp/claude-1000/le_bughunt"); TMP.mkdir(parents=True, exist_ok=True)
HEADS = ["LOG BOOK DATE", "PRO NUMBER", "SITE CODE CANDIDATE", "NAME", "DUE DATE", "AMOUNT", "PAYMENT METHOD", "ACCOUNTANT", "REQUESTOR", "TEAM", "SYSTEM", "COMPANY CODE", "REMARK"]
GROUPS = set(sys.argv[1:]) or set("ABCDEFGHJK")


def row(i, **kw):
    r = dict(zip(HEADS, [datetime.datetime(2026, 10, 1), f"BUG-{i}", "BKK4874A", f"Vendor {i}", datetime.datetime(2026, 10, 15), 1000 + i,
                         "CHEQUE", "Acc", "Req", "IBC", "SITE RENTAL", "1019", None]))
    r.update(kw)
    return r


def xlsx(name, rows, heads=HEADS, sheet="LOG ELECTRIC", blank_after=None):
    wb = Workbook(); ws = wb.active; ws.title = sheet
    ws.append(heads)
    for i, r in enumerate(rows):
        if blank_after is not None and i == blank_after:
            ws.append([])
        ws.append([r.get(h) for h in heads])
    p = TMP / name; wb.save(p); return p


def upload(p, path):
    p.click("#tab-import")
    p.set_input_files("#imp-file", str(path))
    p.wait_for_selector("#imp-ready, #imp-rejected, #imp-noheader, #imp-missing, #imp-empty, #imp-result .card.err", timeout=30000)
    return p.inner_text("#imp-result")


def group_a(b, srv):
    ctx, p = open_page(b, srv)
    # A1 หัวอยู่แถว 3 คอลัมน์สลับ หัวตัวเล็กช่องว่างเกิน มีคอลัมน์เกิน ชีตชื่ออื่น
    shuffled = ["remark", "  team ", "AMOUNT", "pro number", "LOG BOOK DATE", "SITE CODE CANDIDATE", "NAME", "DUE DATE", "PAYMENT METHOD", "ACCOUNTANT", "REQUESTOR", "SYSTEM", "COMPANY CODE", "NOTE ของทีม"]
    data = [None, "IBC", 1200, "BUG-A1", datetime.datetime(2026, 10, 1), "BKK4874A", "V", datetime.datetime(2026, 10, 9), "CHEQUE", "a", "r", "SITE RENTAL", "1019", "x"]
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"; ws.append(["LOG ELECTRIC ตุลาคม"]); ws.append([]); ws.append(shuffled)
    for _ in range(3):
        ws.append(data)
    wb.save(TMP / "a1.xlsx")
    t = upload(p, TMP / "a1.xlsx")
    check("A1 หัวอยู่แถว 3 สลับคอลัมน์ ตัวเล็ก มีคอลัมน์เกิน ชีตชื่อ Sheet1 อ่านได้ 3 แถว", "All 3 rows are valid" in t, t[:150])
    # A2 แถวว่างคั่น เลขแถวในไฟล์ error ต้องตรง Excel
    rows = [row(1), row(2), row(3, AMOUNT="abc"), row(4)]
    with p.expect_download() as d:
        upload(p, xlsx("a2.xlsx", rows, blank_after=2))
    d.value.save_as(TMP / "a2_errors.xlsx")
    ws = load_workbook(TMP / "a2_errors.xlsx").active
    errs = {c[0].row: c[-1].value for c in ws.iter_rows(min_row=2) if c[-1].value}
    check("A2 มีแถวว่างคั่น ไฟล์ error ชี้แถว 5 ตรงกับ Excel (หัวแถว 1 แถวว่างแถว 4)", list(errs) == [5] and str(errs[5]).startswith("Row 5:"), str(errs))
    # A3 วันที่หลายแบบ
    rows = [row(1, **{"LOG BOOK DATE": "2026-10-29"}), row(2, **{"LOG BOOK DATE": "29/10/2569"}), row(3, **{"LOG BOOK DATE": 46324}), row(4, **{"LOG BOOK DATE": "1/2/26"})]
    t = upload(p, xlsx("a3.xlsx", rows))
    check("A3 วันที่ ISO , พ.ศ. , serial ของ Excel , d/m/yy ผ่านทั้งหมด", "All 4 rows are valid" in t, t[:150])
    # A4 ช่องที่เป็นสูตร =NA()
    wb = Workbook(); ws = wb.active; ws.title = "LOG ELECTRIC"; ws.append(HEADS)
    r = row(1); ws.append([r[h] for h in HEADS]); ws["F2"] = "=NA()"
    wb.save(TMP / "a4.xlsx")
    p.click("#tab-import"); p.set_input_files("#imp-file", str(TMP / "a4.xlsx"))
    p.wait_for_selector("#imp-ready, #imp-rejected, #imp-result .card.err", timeout=30000)
    t = p.inner_text("#imp-result")
    check("A4 ช่อง AMOUNT เป็นสูตร =NA() ถูกตีกลับที่ AMOUNT ไม่ผ่านหลอก", "AMOUNT" in t and "rejected" in t, t[:200])
    # A5 มีแต่หัว , A6 ไฟล์ปลอม , A8 ช่องมีแต่ช่องว่าง
    t = upload(p, xlsx("a5.xlsx", []))
    check("A5 ไฟล์มีแต่หัว บอกว่าไม่มีแถวข้อมูล", "no data rows" in t, t[:120])
    (TMP / "a6.xlsx").write_bytes(b"%PDF-1.4 not really excel")
    t = upload(p, TMP / "a6.xlsx")
    check("A6 ไฟล์ปลอม (PDF ตั้งชื่อ .xlsx) ขึ้นข้อความอ่านไม่ได้ ไม่ค้าง", "Cannot read" in t or "not found" in t, t[:150])
    with p.expect_download():
        t = upload(p, xlsx("a8.xlsx", [row(1, NAME="   ", REQUESTOR="\t")]))
    check("A8 ช่องที่มีแต่ช่องว่างนับเป็นว่าง (Required)", "NAME: Required" in t and "REQUESTOR: Required" in t, t[:250])
    # A9 หลายรหัสในช่อง SITE CODE
    t = upload(p, xlsx("a9.xlsx", [row(1, **{"SITE CODE CANDIDATE": "bkk4874a, BKK4875"})]))
    check("A9 สองรหัสคั่นจุลภาคผ่าน", "1 row is valid" in t, t[:120])
    # A10 ไฟล์ส่งออกของแอปอัปโหลดกลับ (มี INDEX CREATED ON ฯลฯ)
    p.click("#tab-records"); p.click("#btn-export")
    with p.expect_download() as d:
        p.click("#export-go")
    exp = TMP / "a10_export.xlsx"; d.value.save_as(exp)
    n = load_workbook(exp).active.max_row - 1
    t = upload(p, exp)
    check("A10 ไฟล์ส่งออกของแอปอัปโหลดกลับ ระบบเตือนว่าจะได้แถวซ้ำ (ไม่ใช่รับเงียบ)", "duplicate" in t.lower() or "already" in t.lower(), f"{n} แถว: {t[:160]}")
    # A11 ไฟล์เทมเพลตที่มีแถวหน้าตาซ้ำกับในตาราง (good.xlsx 40 แถวนำเข้าไปแล้ว) ต้องเตือน
    t = upload(p, pathlib.Path(__file__).parent / "files" / "good.xlsx")
    check("A11 แถวซ้ำกับในตาราง เตือนจำนวนแถวซ้ำ แต่ยังนำเข้าได้", "Everything in this file is already in the table" in t and p.inner_text("#btn-confirm-import").startswith("Import anyway") and p.is_visible("#btn-confirm-import"), t[:200])
    check("A ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
    ctx.close()


def group_b(b, srv):
    ctx, p = open_page(b, srv)
    p.click("#btn-add")
    p.fill("#f-logBookDate", "01/10/2026"); p.fill("#f-proNumber", "BUG-B1"); p.fill("#f-siteCode", "BKK4874A")
    p.fill("#f-name", '<img src=x onerror="window.__xss=1">'); p.fill("#f-dueDate", "31/02/2026"); p.fill("#f-amount", "-150.5")
    p.select_option("#f-paymentMethod", "CHEQUE"); p.fill("#f-accountant", "a"); p.fill("#f-requestor", "r")
    p.select_option("#f-team", "IBC"); p.select_option("#f-system", "DASHBOARD"); p.select_option("#f-companyCode", "1068")
    p.click("#btn-save")
    check("B1 วันที่ 31/02/2026 ขึ้นข้อความใต้ DUE DATE ไม่บันทึก", "Invalid date" in p.inner_text("#e-dueDate") and count(flt="le_pronumber eq 'BUG-B1'") == 0)
    p.fill("#f-dueDate", "15/10/2026")
    p.click("#btn-save")
    p.wait_for_function("() => (document.querySelector('#toast')?.innerText || '').startsWith('Added')", timeout=30000)
    p.fill("#rec-search", "BUG-B1"); p.wait_for_timeout(500)
    check("B2 ชื่อที่มีโค้ด HTML แสดงเป็นข้อความ ไม่ทำงาน", p.evaluate("!window.__xss && !document.querySelector('#rec-table img')"))
    check("B2 AMOUNT ติดลบบันทึกได้", "-150.50" in p.inner_text("#rec-table"))
    # B3 ลบ COMPANY CODE ตอนแก้ COMPANY NAME ต้องว่างตาม
    p.click("#rec-table tbody tr[data-id]"); p.click("#btn-edit"); p.select_option("#f-companyCode", ""); p.click("#btn-save")
    p.wait_for_selector("#hist-list .diff")
    rec = dv("GET", "le_logelectrics?$select=le_companycode,le_companyname,le_logelectricid&$filter=le_pronumber eq 'BUG-B1'")["value"][0]
    check("B3 ลบ COMPANY CODE แล้ว COMPANY NAME ว่างตาม และประวัติบอกทั้งสองช่อง", not rec["le_companycode"] and not rec["le_companyname"] and "COMPANY NAME" in p.inner_text("#hist-list"), json.dumps(rec)[:150])
    # B4 สองคนแก้พร้อมกัน: เปิดฟอร์มค้างไว้ อีกคนแก้ผ่าน API ก่อน แล้วค่อยกดบันทึก
    p.click("#btn-edit")
    dv("PATCH", f"le_logelectrics({rec['le_logelectricid']})", {"le_remark": "แก้จากอีกเครื่อง"})
    p.evaluate("document.querySelector('#toast').innerText = ''")
    p.fill("#f-accountant", "changed in form"); p.click("#btn-save")
    p.wait_for_function("() => (document.querySelector('#toast')?.innerText || '').length > 0", timeout=30000)
    t = p.inner_text("#toast")
    after = dv("GET", f"le_logelectrics({rec['le_logelectricid']})?$select=le_remark,le_accountant")
    check("B4 แก้ชนกัน แอปบอกว่ามีคนแก้ก่อน และไม่ทับค่าของอีกคน", "Someone else changed" in t and after["le_remark"] == "แก้จากอีกเครื่อง", f"{t} / {after}")
    check("B ไม่มี error ในคอนโซล (ยกเว้น 412 ที่ตั้งใจ)", not [e for e in p.errors if "412" not in e], "; ".join(p.errors[:3]))
    ctx.close()


def group_c(b, srv):
    # ROLE_REMOVE_APPROVED: เป้าหมายคือบัญชี 64 (ไม่ใช่ 65 ที่ผูก az) จด baseline ลงไฟล์ก่อนแตะ ตอนจบถอดเฉพาะบทบาทที่ไม่มีใน baseline
    u64 = user_id("64050580")
    base = sorted(r["name"] for r in dv("GET", f"systemusers({u64})/systemuserroles_association?$select=name")["value"])
    (EVID / "bughunt_roles_baseline_64.json").write_text(json.dumps(base, ensure_ascii=False))
    mgr = dv("GET", "roles?$select=roleid,name&$filter=" + urllib.parse.quote("name eq 'LE Manager'"))["value"][0]["roleid"]
    # C5 (รันก่อนใส่บทบาท) Member แอบยิงแก้แถวของ 65 และแอบเขียนประวัติให้แถวคนอื่น
    # Dataverse จำสิทธิ์เดิมราว 1 นาทีหลังเปลี่ยนบทบาท (วัด 29/10: ถอด Manager แล้ว 20:22:42 ยังแก้ได้ 20:23:45 โดนปฏิเสธ) จึงต้องไม่รันต่อจากการถอดบทบาท
    rid = dv("GET", "le_logelectrics?$select=le_logelectricid&$filter=le_pronumber eq 'PRO100006'")["value"][0]["le_logelectricid"]
    codes = []
    for m, path, body in (("PATCH", f"le_logelectrics({rid})", {"le_amount": 1}),
                          ("POST", "le_histories", {"le_name": "fake", "le_record@odata.bind": f"/le_logelectrics({rid})"})):
        try:
            dv(m, path, body, caller=u64); codes.append(200)
        except urllib.error.HTTPError as e:
            codes.append(e.code)
    check("C5 Member แก้แถวคนอื่นและเขียนประวัติปลอมให้แถวคนอื่นไม่ได้", all(c >= 400 for c in codes), str(codes))
    added = "LE Manager" not in base
    if added:
        dv("POST", f"systemusers({u64})/systemuserroles_association/$ref", {"@odata.id": f"{BASE}/api/data/v9.2/roles({mgr})"})
    t0 = time.time()
    for _ in range(120):  # รอจนสิทธิ์ Manager มีผลจริงก่อนเปิดหน้า (บางรอบ Dataverse ใช้เกิน 2 นาที เห็นจริง 29/10 , 30/10 เกิน 4 นาที)
        if dv("GET", "le_logelectrics?$select=le_logelectricid&$count=true&$top=1", caller=u64)["@odata.count"] == count():
            break
        time.sleep(5)
    # หมดเวลารอแล้วสิทธิ์ยังไม่มีผล ฟ้องตรง ๆ ว่าเป็นความหน่วงของ Dataverse ไม่ใช่ปล่อยไปแดงที่ C1 เหมือนแอปผิด
    check("C Dataverse ให้สิทธิ์ Manager มีผลภายใน 10 นาทีหลังใส่บทบาท", time.time() - t0 < 600, f"รอ {time.time() - t0:.0f} วินาที")
    try:
        ctx, p = open_page(b, srv, caller=u64)
        check("C1 64 เป็น Manager เห็นป้าย Manager ไม่มีแท็บ Settings", p.inner_text("#my-role") == "Manager" and not p.query_selector("#tab-settings"), p.inner_text("#my-role"))
        n_all = count()
        check("C1 Manager เห็นทุกแถว", p.inner_text("#rec-summary").startswith(f"{n_all} of"), f"{p.inner_text('#rec-summary')} / {n_all}")
        p.fill("#rec-search", "PRO100005"); p.wait_for_timeout(400)
        p.click("#rec-table tbody tr[data-id]")
        check("C2 Manager ไม่มีปุ่มลบแถว", not p.query_selector("#btn-del-one"))
        p.click("#btn-edit"); p.fill("#f-remark", "manager edit"); p.click("#btn-save")
        p.wait_for_selector("#hist-list .diff")
        rid = dv("GET", "le_logelectrics?$select=le_logelectricid,_ownerid_value&$filter=le_pronumber eq 'PRO100005'")["value"][0]
        h = dv("GET", f"le_histories?$select=_ownerid_value,_createdby_value&$filter=_le_record_value eq {rid['le_logelectricid']}&$orderby=createdon desc&$top=1")["value"][0]
        check("C3 Manager แก้แถวของ Admin ประวัติเป็นของเจ้าของแถว (เจ้าของอ่านได้) คนสร้างคือ Manager", h["_ownerid_value"] == rid["_ownerid_value"] and h["_createdby_value"] == u64, json.dumps(h)[:200])
        try:
            dv("DELETE", f"le_logelectrics({rid['le_logelectricid']})", caller=u64); ok = False
        except urllib.error.HTTPError as e:
            ok = e.code == 403
        check("C4 Manager ยิงลบตรงไม่ได้ (403)", ok)
        check("C ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
        ctx.close()
    finally:
        if added:  # ROLE_REMOVE_APPROVED ถอดเฉพาะ LE Manager ที่เทสใส่เอง
            req = urllib.request.Request(f"{BASE}/api/data/v9.2/systemusers({u64})/systemuserroles_association/$ref?$id={BASE}/api/data/v9.2/roles({mgr})", method="DELETE",
                                         headers={"Authorization": f"Bearer {os.environ['PVDESK_TOKEN']}"})
            urllib.request.urlopen(req).read()
    now = sorted(r["name"] for r in dv("GET", f"systemusers({u64})/systemuserroles_association?$select=name")["value"])
    check("C บทบาทของ 64 กลับเท่า baseline", now == base, f"{now} / {base}")


def group_d(b, srv):
    ctx, p = open_page(b, srv)
    p.click("#btn-add")
    vals = {"#f-logBookDate": "01/10/2026", "#f-proNumber": "BUG-D1", "#f-siteCode": "BKK4874A", "#f-name": "=HYPERLINK(\"http://x\",\"click\")",
            "#f-dueDate": "15/10/2026", "#f-amount": "10", "#f-accountant": "+cmd", "#f-requestor": "@SUM(1)"}
    for k, v in vals.items():
        p.fill(k, v)
    p.select_option("#f-paymentMethod", "CHEQUE"); p.select_option("#f-team", "IBC"); p.select_option("#f-system", "DASHBOARD")
    p.click("#btn-save")
    p.wait_for_function("() => (document.querySelector('#toast')?.innerText || '').startsWith('Added')", timeout=30000)
    p.fill("#rec-search", "BUG-D1"); p.wait_for_timeout(400)
    p.click("#btn-export"); p.check('#export-dlg input[data-fmt="csv"]'); p.uncheck('#export-dlg input[data-fmt="xlsx"]')
    with p.expect_download() as d:
        p.click("#export-go")
    rows = list(csv.reader(io.StringIO(pathlib.Path(d.value.path()).read_text(encoding="utf-8-sig"))))
    hd = rows[0]; r = rows[1]
    cells = [r[hd.index("NAME")], r[hd.index("ACCOUNTANT")], r[hd.index("REQUESTOR")]]
    check("D1 CSV กันสูตร: ค่าที่ขึ้นต้น = + @ ไม่ถูก Excel ตีเป็นสูตร", not any(c[:1] in "=+-@" for c in cells), str(cells))
    ctx.close()


def group_e(b, srv):
    N = 30000
    fake = []
    for i in range(N):
        fake.append({"le_logelectricid": f"00000000-0000-0000-0000-{i:012d}", "@odata.etag": 'W/"1"', "le_index": str(i + 1), "le_logbookdate": f"2026-{i % 12 + 1:02d}-{i % 28 + 1:02d}",
                     "le_pronumber": f"PRO{i:06d}", "le_sitecode": f"BKK{i % 9000:04d}A", "le_sitecodenocand": f"BKK{i % 9000:04d}", "le_name": f"ผู้ให้เช่า {i % 700}",
                     "le_duedate": "2026-11-13", "le_amount": 1000 + i % 5000, "le_paymentmethod": ["CHEQUE", "BANK TRANSFER"][i % 2], "le_accountant": f"Acc {i % 9}",
                     "le_requestor": f"Req {i % 40}", "le_team": ["IBC", "MACRO", "ONLINE"][i % 3], "le_system": ["SITE RENTAL", "DASHBOARD"][i % 2], "le_companycode": "1068",
                     "le_companyname": "TRUE MOVE H UNIVERSAL COMMUNICATION CO., LTD.", "le_remark": None, "le_source": "fake.xlsx", "createdon": "2026-10-29T10:00:00Z",
                     "modifiedon": "2026-10-29T10:00:00Z", "_createdby_value": "u", "_createdby_value@OData.Community.Display.V1.FormattedValue": "Fake User",
                     "_modifiedby_value": "u", "_modifiedby_value@OData.Community.Display.V1.FormattedValue": "Fake User", "_ownerid_value": "u"})
    body = json.dumps({"value": fake})
    ctx = b.new_context(viewport={"width": 1440, "height": 900}, accept_downloads=True)
    p = ctx.new_page(); p.errors = []
    p.on("pageerror", lambda e: p.errors.append(str(e)))
    local = f"http://127.0.0.1:{srv.port}"

    def fwd(route, request):  # ข้อมูลจำลองล้วน ไม่แตะ Dataverse
        u = urllib.parse.unquote(request.url)
        if "/le_logelectrics?$select" in u:
            return route.fulfill(status=200, content_type="application/json", body=body)
        if u.endswith("/WhoAmI"):
            return route.fulfill(status=200, content_type="application/json", body='{"UserId":"u"}')
        if "/systemusers(u)" in u:
            return route.fulfill(status=200, content_type="application/json", body='{"fullname":"Fake Admin","value":[{"name":"LE Admin"}]}')
        return route.fulfill(status=200, content_type="application/json", body='{"value":[]}')
    p.route(f"{local}/api/data/v9.2/**", fwd)
    t0 = time.time(); p.goto(srv.url); p.wait_for_selector("body.app-ready", timeout=60000); boot = time.time() - t0
    check(f"E1 เปิดแอปกับ {N:,} แถว ภายใน 6 วินาที", boot < 6 and p.inner_text("#rec-summary").startswith("30,000 of"), f"{boot:.2f}s {p.inner_text('#rec-summary')[:40]}")
    t0 = time.time(); p.fill("#rec-search", "ผู้ให้เช่า 15"); p.wait_for_function("() => !document.querySelector('#rec-summary').innerText.startsWith('30,000')"); s1 = time.time() - t0 - 0.2
    check("E2 ค้นหาทุกคอลัมน์ใน 30,000 แถว ตอบภายใน 1 วินาที (หลังหน่วงพิมพ์ 0.2)", s1 < 1, f"{s1:.2f}s")
    p.fill("#rec-search", ""); p.wait_for_timeout(600)
    p.hover('#rec-table th[data-col="requestor"]')
    t0 = time.time(); p.click('#rec-table th[data-col="requestor"] .xf-btn'); p.wait_for_selector(".xf-pop .xf-item"); s2 = time.time() - t0
    check("E3 เปิดตัวกรอง REQUESTOR ภายใน 1.5 วินาที", s2 < 1.5, f"{s2:.2f}s")
    p.keyboard.press("Escape")
    t0 = time.time(); p.select_option("#rec-group", "requestor"); p.wait_for_selector("#rec-table tr.grp"); s3 = time.time() - t0
    check("E4 Group by REQUESTOR ภายใน 1.5 วินาที", s3 < 1.5, f"{s3:.2f}s")
    p.select_option("#rec-group", "")
    p.click("#btn-export")
    t0 = time.time()
    with p.expect_download(timeout=120000) as d:
        p.click("#export-go")
    s4 = time.time() - t0
    check("E5 ส่งออก XLSX 30,000 แถว ภายใน 20 วินาที", s4 < 20, f"{s4:.2f}s {os.path.getsize(d.value.path()) // 1024} KB")
    check("E ไม่มี error ในหน้า", not p.errors, "; ".join(p.errors[:3]))
    ctx.close()


def group_f(b, srv):
    ctx, p = open_page(b, srv, vp=(390, 844))
    p.click("#btn-add")
    p.wait_for_selector("#rec-form")
    p.wait_for_timeout(400)
    box = p.evaluate("""() => { const s = document.querySelector('#btn-save').getBoundingClientRect();
      return { saveTop: Math.round(s.top), saveBottom: Math.round(s.bottom), h: innerHeight, overflow: document.documentElement.scrollWidth > innerWidth + 1 }; }""")
    check("F1 มือถือ: ฟอร์มเพิ่มแถวเห็นปุ่ม Save และไม่ล้นแนวนอน", box["saveBottom"] <= box["h"] and box["saveTop"] >= 0 and not box["overflow"], str(box))
    p.screenshot(path=str(EVID / "bug_f1_mobile_form.png"))
    check("F ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
    ctx.close()


def fill_of(ws, ref):
    c = ws[ref]
    return c.fill.fgColor.rgb if c.fill and c.fill.fill_type == "solid" else None


def dl(d):
    """บันทึกไฟล์ที่ดาวน์โหลดตามชื่อจริง (path ชั่วคราวของ Playwright ไม่มีนามสกุล openpyxl ไม่ยอมเปิด)"""
    out = TMP / d.value.suggested_filename; d.value.save_as(out); return out


def view(p, r):
    p.click("#my-role"); p.click(f'#role-menu [data-role="{r}"]'); p.wait_for_timeout(300)


def group_g(b, srv):
    """รอบ 2 (พี่ปอนด์ 29/10 ดึก "ดูบัคอีกรอบ") ของที่เพิ่มหลังรอบแรก: สลับบทบาท , ลากขยายคอลัมน์ , แถวตัวอย่าง , ไฟล์ error ลงสี , การ์ดฉลอง"""
    from openpyxl.utils import get_column_letter as gcl
    u65 = user_id("65050782")
    ctx, p = open_page(b, srv)
    n_all, n_me = count(), count(flt=f"_ownerid_value eq {u65}")
    # G1 ดูแบบ Member แล้วเพิ่มแถว แถวใหม่ต้องขึ้นในมุม Member ทันที (เจ้าของคือตัวเอง)
    view(p, "MEMBER")
    p.click("#btn-add")
    for k, v in {"#f-logBookDate": "01/10/2026", "#f-proNumber": "BUG-G1", "#f-siteCode": "BKK4874A", "#f-name": "G", "#f-dueDate": "15/10/2026",
                 "#f-amount": "10", "#f-accountant": "a", "#f-requestor": "r"}.items():
        p.fill(k, v)
    p.select_option("#f-paymentMethod", "CHEQUE"); p.select_option("#f-team", "IBC"); p.select_option("#f-system", "DASHBOARD")
    p.click("#btn-save")
    p.wait_for_function("() => (document.querySelector('#toast')?.innerText || '').startsWith('Added')", timeout=30000)
    p.wait_for_timeout(500)
    check("G1 ดูแบบ Member แล้วเพิ่มแถว แถวใหม่ขึ้นในมุม Member และยังเห็นเฉพาะแถวของตัวเอง", p.inner_text("#rec-summary").startswith(f"{n_me + 1} of")
          and p.inner_text("#my-role") == "Member", f"{p.inner_text('#rec-summary')} / ของตัวเอง {n_me} + 1")
    # G2 ดูแบบ Member แล้วส่งออก ได้เฉพาะแถวของตัวเอง
    p.click("#btn-export")
    txt = p.inner_text("#export-dlg")
    with p.expect_download() as d:
        p.click("#export-go")
    got = load_workbook(dl(d)).active.max_row - 1
    check("G2 ดูแบบ Member ส่งออกได้เฉพาะแถวของตัวเอง และบอกว่า your rows only", got == n_me + 1 and "your rows only" in txt, f"{got} / {n_me + 1}")
    # G3 ดูแบบ Member แล้วกดรีเฟรช ยังกรองเฉพาะของตัวเอง
    p.click("#btn-refresh") if p.query_selector("#btn-refresh") else p.click('button[title*="Refresh"], button[aria-label*="Refresh"]')
    p.wait_for_timeout(2500)
    check("G3 ดูแบบ Member แล้วกดรีเฟรช ยังเห็นเฉพาะแถวของตัวเอง", p.inner_text("#rec-summary").startswith(f"{n_me + 1} of"), p.inner_text("#rec-summary"))
    view(p, "ADMIN")
    check("G4 กลับเป็น Admin เห็นทุกแถวรวมแถวใหม่", p.inner_text("#rec-summary").startswith(f"{n_all + 1} of"), p.inner_text("#rec-summary"))
    # G5 ลากขอบหดจนสุด ต้องหยุดที่ 48px ไม่หายไป
    loc = p.locator('#rec-table th[data-col="name"] .col-rs'); loc.scroll_into_view_if_needed(); g = loc.bounding_box()
    x, y = g["x"] + g["width"] / 2, g["y"] + g["height"] / 2
    p.mouse.move(x, y); p.mouse.down(); p.mouse.move(x - 300, y); p.mouse.move(x - 600, y); p.mouse.up(); p.wait_for_timeout(300)
    w = p.evaluate("Math.round(document.querySelector('#rec-table th[data-col=\"name\"]').getBoundingClientRect().width)")
    check("G5 ลากขอบหดสุด คอลัมน์หยุดที่ราว 48px ไม่หายไป", 44 <= w <= 52, str(w))
    # G6 ความกว้างที่ลากไว้ยังอยู่ตอน Group by
    p.select_option("#rec-group", "system"); p.wait_for_timeout(400)
    w2 = p.evaluate("Math.round(document.querySelector('#rec-table th[data-col=\"name\"]').getBoundingClientRect().width)")
    check("G6 เปิด Group by แล้วความกว้างที่ลากไว้ยังอยู่", abs(w2 - w) <= 2, f"{w} -> {w2}")
    p.select_option("#rec-group", ""); p.evaluate("localStorage.removeItem('le.widths')")
    # G7 ไฟล์มีแถวตัวอย่างปนแถวผิด ตีกลับเฉพาะแถวผิด แถวตัวอย่างไม่มี ERROR
    ex = {"LOG BOOK DATE": datetime.datetime(2026, 10, 29), "PRO NUMBER": "PRO Group", "SITE CODE CANDIDATE": "BKK4874A", "NAME": "ABC Property Co., Ltd.",
          "DUE DATE": datetime.datetime(2026, 11, 13), "AMOUNT": 2021.54, "PAYMENT METHOD": "BANK TRANSFER", "ACCOUNTANT": "Patcharasit Pongudom",
          "REQUESTOR": "Jirachaya Rachpolsaen", "TEAM": "IBC", "SYSTEM": "SITE RENTAL", "COMPANY CODE": "1068", "REMARK": "Meter 5.13 kW"}
    with p.expect_download() as d:
        upload(p, xlsx("g7.xlsx", [ex, row(1, AMOUNT="x")]))
    ws = load_workbook(dl(d)).active
    ec = [c.value for c in ws[1]].index("ERROR") + 1
    check("G7 แถวตัวอย่างปนแถวผิด ไฟล์ error ชี้แค่แถว 3 แถวตัวอย่างไม่มี ERROR", ws.cell(2, ec).value is None and str(ws.cell(3, ec).value).startswith("Row 3"),
          f"{ws.cell(2, ec).value} | {ws.cell(3, ec).value}")
    # G8 แถวตัวอย่างใน CSV ก็ถูกข้าม
    pth = TMP / "g8.csv"
    with open(pth, "w", newline="", encoding="utf-8-sig") as fh:
        wr = csv.writer(fh); wr.writerow(HEADS)
        for r in (ex, row(2)):
            wr.writerow([v.strftime("%d/%m/%Y") if isinstance(v, datetime.datetime) else ("" if v is None else v) for v in (r[h] for h in HEADS)])
    t = upload(p, pth)
    check("G8 CSV ที่มีแถวตัวอย่าง ข้ามแถวตัวอย่าง นำเข้าแถวจริง 1 แถว", "1 row is valid" in t and p.is_visible("#imp-example"), t[:200])
    # G9 ไฟล์ error ของไฟล์หัวอยู่แถว 3 คอลัมน์สลับ มีคอลัมน์เกิน: ลงสีถูกคอลัมน์ ถูกแถว
    shuffled = ["remark", "  team ", "AMOUNT", "pro number", "LOG BOOK DATE", "SITE CODE CANDIDATE", "NAME", "DUE DATE", "PAYMENT METHOD", "ACCOUNTANT", "REQUESTOR", "SYSTEM", "COMPANY CODE", "NOTE ของทีม"]
    data = [None, "IBC", "abc", "BUG-G9", datetime.datetime(2026, 10, 1), "BKK4874A", "V", datetime.datetime(2026, 10, 9), "CHEQUE", "a", "r", "SITE RENTAL", "1019", "x"]
    wb = Workbook(); w0 = wb.active; w0.title = "Sheet1"; w0.append(["LOG ELECTRIC ตุลาคม"]); w0.append([]); w0.append(shuffled); w0.append(data)
    wb.save(TMP / "g9.xlsx")
    with p.expect_download() as d:
        upload(p, TMP / "g9.xlsx")
    ws = load_workbook(dl(d)).active
    hdr = [c.value for c in ws[3]]
    E = gcl(hdr.index("ERROR") + 1)
    check("G9 หัวอยู่แถว 3: หัว ERROR แดงที่แถว 3 ช่อง AMOUNT (คอลัมน์ C) ของแถว 4 แดงอ่อน คอลัมน์เกินไม่ลงสีหัว", fill_of(ws, f"{E}3") == "FFC00000" and fill_of(ws, "C4") == "FFFDECEC"
          and fill_of(ws, "A3") == "FFF2F2F2" and fill_of(ws, "C3") == "FFFDECEC" and fill_of(ws, "N3") is None and fill_of(ws, "D4") is None,
          f"E3 {fill_of(ws, E + '3')} C4 {fill_of(ws, 'C4')} A3 {fill_of(ws, 'A3')} N3 {fill_of(ws, 'N3')}")
    # G10 คอลัมน์เกิน Z (30 คอลัมน์) ไฟล์ error ยังลงสีถูกและเปิดได้
    heads30 = HEADS + [f"X{i}" for i in range(17)]
    r30 = row(3, TEAM="DTAC"); r30.update({f"X{i}": i for i in range(17)})
    with p.expect_download() as d:
        upload(p, xlsx("g10.xlsx", [r30], heads=heads30))
    ws = load_workbook(dl(d)).active
    hdr = [c.value for c in ws[1]]
    E = gcl(hdr.index("ERROR") + 1)
    check("G10 ไฟล์ 30 คอลัมน์ ERROR อยู่คอลัมน์ AE ลงสีครบ ช่อง TEAM แดงอ่อน", E == "AE" and fill_of(ws, "AE1") == "FFC00000" and fill_of(ws, "AE2") == "FFFFF0F0" and fill_of(ws, "J2") == "FFFDECEC",
          f"{E} {fill_of(ws, 'AE1')} {fill_of(ws, 'J2')}")
    # G11 ไฟล์ error ของ CSV ก็ลงสี
    pth = TMP / "g11.csv"
    with open(pth, "w", newline="", encoding="utf-8-sig") as fh:
        wr = csv.writer(fh); wr.writerow(HEADS); r = row(4, AMOUNT="สองพัน")
        wr.writerow([v.strftime("%d/%m/%Y") if isinstance(v, datetime.datetime) else ("" if v is None else v) for v in (r[h] for h in HEADS)])
    with p.expect_download() as d:
        upload(p, pth)
    ws = load_workbook(dl(d)).active
    check("G11 ไฟล์ error ของ CSV ลงสี หัว ERROR แดง ช่อง AMOUNT แดงอ่อน", fill_of(ws, "N1") == "FFC00000" and fill_of(ws, "F2") == "FFFDECEC", f"{fill_of(ws, 'N1')} {fill_of(ws, 'F2')}")
    # G12 นำเข้า 1 แถว ข้อความเอกพจน์ และนำเข้าซ้ำ การ์ดฉลองมีใบเดียว
    for i in (1, 2):
        upload(p, xlsx(f"g12_{i}.xlsx", [row(20 + i)]))
        p.click("#btn-confirm-import"); p.click("dialog.dlg[open] .btn.dark")  # หน้าต่าง Export ที่ปิดแล้วยังอยู่ใน DOM เลือกเฉพาะที่เปิดอยู่
        p.wait_for_selector("#imp-done", timeout=60000)
    t = p.inner_text("#imp-done .done-title")
    check("G12 นำเข้า 1 แถวเขียน Imported 1 row และนำเข้าซ้ำแล้วการ์ดฉลองเหลือใบเดียว", t == "Imported 1 row" and p.locator("#imp-done").count() == 1
          and p.inner_text("#toast .toast-msg") == "Imported 1 row", t)
    # G16 ข้อความแถวซ้ำต้องบอกว่ากี่แถวจากทั้งหมด และแถวไหนบ้าง ภาษาถูก (พี่ปอนด์ถาม 29/10 ว่าข้อความนี้แปลว่าซ้ำทุกแถวไหม)
    ex2 = []
    for pro in ("PRO100000", "PRO100001"):
        v = dv("GET", "le_logelectrics?$select=le_logbookdate,le_sitecode,le_amount&$filter=" + urllib.parse.quote(f"le_pronumber eq '{pro}'"))["value"][0]
        y, m, d_ = map(int, v["le_logbookdate"][:10].split("-"))
        ex2.append(row(30, **{"LOG BOOK DATE": datetime.datetime(y, m, d_), "PRO NUMBER": pro, "SITE CODE CANDIDATE": v["le_sitecode"], "AMOUNT": v["le_amount"]}))
    t = upload(p, xlsx("g16.xlsx", ex2 + [row(31)]))
    check("G16 ไฟล์ 3 แถวซ้ำ 2 แถว ข้อความบอก 2 of 3 rows are already in the table และ Rows 2, 3", "2 of 3 rows are already in the table" in t and "Rows 2, 3" in t, t[:400])
    # ซ้ำบางแถวก็ขึ้นการ์ดเตือนแบบเดียวกับซ้ำทุกแถว ไม่ใช่ข้อความบรรทัดเล็ก , บรรทัดชีตบอกจำนวนแถว ไม่ใช่ rows 2 to 4 (พี่ปอนด์ 30/10 ภาพ #77 #78 #80)
    check("G16 ซ้ำบางแถวขึ้นการ์ดเตือน (หัวตัวหนา) ปุ่มนำเข้ายังเป็นปุ่มหลัก บรรทัดชีตบอก 3 rows to import",
          p.is_visible("#imp-dup.alldup b") and p.inner_text("#imp-dup b") == "2 of 3 rows are already in the table"
          and "btn primary" in p.get_attribute("#btn-confirm-import", "class") and "Sheet LOG ELECTRIC, 3 rows to import." in t and " to 4" not in t, t[:300])
    t = upload(p, xlsx("g16b.xlsx", ex2[:1]))
    check("G16 ไฟล์แถวเดียวที่ซ้ำ (ซ้ำทุกแถว) ขึ้นการ์ดเตือนแยก ไม่มีคำว่า 1 rows ปุ่ม Import anyway", p.is_visible("#imp-alldup") and "1 rows" not in t
          and p.inner_text("#btn-confirm-import") == "Import anyway 1 row", t[:400])
    # G18 Admin กับ Manager เลือกชื่อคนเพื่อดูแถวของคนนั้น (พี่ปอนด์ 29/10 ภาพ #76)
    u64 = user_id("64050580")
    n64 = count(flt=f"_ownerid_value eq {u64}")
    p.click("#tab-records"); p.fill("#rec-search", ""); p.wait_for_timeout(300)
    names = owner_names(p)
    pick_owner(p, u64)
    summ = p.inner_text("#rec-summary")
    check("G18 Admin มีช่องเลือกชื่อคน เลือก 64 แล้วเห็นเฉพาะแถวของ 64 เท่ากับใน Dataverse", bool(names) and names[0].startswith("All people") and summ.startswith(f"{n64} of"), f"{names[:4]} | {summ} | 64 มี {n64}")
    # G19 พิมพ์ค้นหาชื่อในช่องเลือกคน (พี่ปอนด์ 30/10 ภาพ #81) พิมพ์ตัวเล็กบางส่วนของชื่อก็เจอ กด Enter เลือกคนแรก ไม่เจอบอก No people found
    pick_owner(p, "")
    p.click("#rec-owner .owner-btn"); p.fill("#rec-owner-q", "patch"); p.wait_for_timeout(200)
    vis = p.eval_on_selector_all("#rec-owner .owner-opt", "os => os.filter(o => o.offsetParent).map(o => o.textContent)")
    p.keyboard.press("Enter"); p.wait_for_timeout(300)
    summ2, label = p.inner_text("#rec-summary"), p.inner_text("#rec-owner .owner-btn")
    check("G19 พิมพ์ patch ในช่องค้นหาชื่อ เหลือแค่ Patcharasit กด Enter แล้วเลือกให้", len(vis) == 1 and vis[0].startswith("Patcharasit") and summ2.startswith(f"{n64} of")
          and label.startswith("Patcharasit") and p.is_hidden("#rec-owner .owner-pop"), f"{vis} | {label} | {summ2}")
    p.click("#rec-owner .owner-btn"); p.fill("#rec-owner-q", "zzzz"); p.wait_for_timeout(200)
    check("G19 พิมพ์ชื่อที่ไม่มี บอก No people found", "No people found" in p.inner_text("#rec-owner .owner-pop"))
    p.keyboard.press("Escape"); p.wait_for_timeout(200)
    check("G19 กด Esc ปิดรายการ ค่าที่เลือกไว้ยังอยู่", p.is_hidden("#rec-owner .owner-pop") and p.inner_text("#rec-owner .owner-btn").startswith("Patcharasit"))
    pick_owner(p, "")
    view(p, "MEMBER")
    check("G18 ดูแบบ Member ไม่มีช่องเลือกชื่อคน", not p.is_visible("#rec-owner"))
    view(p, "ADMIN")
    # G17 ตัวกรองคอลัมน์วันที่เป็นต้นไม้ ปี > เดือน (อังกฤษ) > วัน แบบ Excel (พี่ปอนด์ 29/10 รอบ 4 ภาพ #65)
    p.click("#tab-records"); p.fill("#rec-search", ""); p.wait_for_timeout(300)
    before = int(p.inner_text("#rec-summary").split(" of ")[0].replace(",", ""))
    p.click('#rec-table th[data-col="logBookDate"] .xf-btn', force=True)
    years = p.eval_on_selector_all('.xf-pop [data-node^="y:"]', "els => els.map(e => e.dataset.node)")
    months = p.eval_on_selector_all('.xf-pop .xf-tree', "els => els.filter(e => e.querySelector('[data-node^=\"m:\"]')).map(e => e.querySelector('.xf-v').innerText)")
    check("G17 ตัวกรอง LOG BOOK DATE เป็นต้นไม้ ปีเรียงน้อยไปมาก กางอยู่ เดือนเป็นชื่ออังกฤษ ยังไม่เห็นวัน", "y:2026" in years and years == sorted(years) and len(months) >= 3
          and all(m in ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December") for m in months)
          and not p.query_selector('.xf-pop [data-node^="d:"]'), f"{years} {months}")
    mnode = p.eval_on_selector('.xf-pop [data-node^="m:"]', "e => e.dataset.node")
    p.click(f'.xf-pop [data-tog="{mnode}"]')
    days = p.eval_on_selector_all(f'.xf-pop [data-node^="d:{mnode[2:]}"]', "els => els.length")
    mcount = int(p.eval_on_selector(f'.xf-pop [data-node="{mnode}"]', "e => e.closest('.xf-tree').querySelector('.xf-n').innerText").replace(",", ""))
    check("G17 กดกางเดือน เห็นวันในเดือนนั้น", days >= 1, f"{mnode} {days} วัน")
    p.uncheck(f'.xf-pop [data-node="{mnode}"]')
    yst = p.eval_on_selector(f'.xf-pop [data-node="y:{mnode[2:6]}"]', "e => e.indeterminate || !e.checked")  # ปีที่มีเดือนเดียวจะเป็นไม่ติ๊กทั้งปี
    p.click(".xf-pop .xf-ok"); p.wait_for_timeout(400)
    after = int(p.inner_text("#rec-summary").split(" of ")[0].replace(",", ""))
    check("G17 เอาติ๊กเดือนออก ปีเป็นติ๊กครึ่ง กด OK แล้วแถวลดเท่าจำนวนของเดือนนั้น", yst and after == before - mcount, f"{before} - {mcount} = {after}")
    p.screenshot(path=str(EVID / "bug_g17_date_tree.png"))
    p.click('#rec-table th[data-col="logBookDate"] .xf-btn', force=True); p.fill(".xf-pop .xf-search", "/10/")
    flat = p.eval_on_selector_all(".xf-pop .xf-list input[data-v]", "els => els.length")
    check("G17 พิมพ์ค้นหาในตัวกรองวันที่ กลับเป็นรายการเรียบให้ติ๊กทีละวัน", flat >= 1 and not p.query_selector(".xf-pop .xf-tree"), str(flat))
    p.keyboard.press("Escape"); p.click("#rec-clear") if p.is_visible("#rec-clear") else None
    p.wait_for_load_state("networkidle")  # หลังนำเข้าแอปโหลดข้อมูลใหม่ ปิดก่อนเสร็จ Route.fetch ล้มไปโผล่ตอนเปิดหน้าถัดไป
    check("G ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
    ctx.close()
    # G13 มือถือ Admin: ไม่ล้นแนวนอน , G15 เห็นป้ายบทบาทและสลับได้
    ctx, p = open_page(b, srv, vp=(390, 844))
    check("G13 มือถือ Admin ไม่ล้นแนวนอน", p.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"))
    p.screenshot(path=str(EVID / "bug_g13_mobile_admin.png"))
    vis = p.is_visible("#my-role")
    if vis:
        view(p, "MEMBER")
    check("G15 มือถือ Admin เห็นป้ายบทบาท กดสลับเป็น Member ได้ และยังไม่ล้นแนวนอน", vis and p.inner_text("#my-role") == "Member"
          and p.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"), f"เห็นป้าย {vis}")
    p.screenshot(path=str(EVID / "bug_g15_mobile_role.png"))
    ctx.close()
    # G14 ตารางเรียง INDEX จากมากไปน้อยต่อเนื่อง แม้แถวจากไฟล์เดียวกันมีเวลาสร้างตรงกันถึงวินาที (Dataverse ส่งลำดับมามั่ว 3..40 แล้ว 1 2)
    ctx, p = open_page(b, srv)
    idx = [int(x) for x in p.eval_on_selector_all('#rec-table tbody td[data-c="index"]', "els => els.map(e => e.innerText.trim())") if x.strip()]
    check("G14 ตารางเรียง INDEX จากมากไปน้อยต่อเนื่อง (ใหม่สุดอยู่บน)", len(idx) > 5 and idx == sorted(idx, reverse=True), str(idx[:8]) + " ... " + str(idx[-4:]))
    ctx.close()


def owner_names(p):
    """ชื่อในช่องเลือกคน (เปิดรายการ อ่าน แล้วปิด)"""
    p.click("#rec-owner .owner-btn"); p.wait_for_selector("#rec-owner .owner-pop", state="visible")
    names = p.eval_on_selector_all("#rec-owner .owner-opt", "os => os.map(o => o.textContent)")
    p.keyboard.press("Escape"); p.wait_for_timeout(150)
    return names


def pick_owner(p, uid):
    """เลือกคนในช่อง All people ("" = ทุกคน)"""
    p.click("#rec-owner .owner-btn"); p.click(f'#rec-owner .owner-opt[data-id="{uid}"]'); p.wait_for_timeout(300)


def toast(p, timeout=30000):
    p.wait_for_function("() => document.querySelector('#toast')?.classList.contains('show')", timeout=timeout)
    return p.inner_text("#toast .toast-msg")


def save(p):
    """ล้างกล่องแจ้งเก่าก่อนกดบันทึก ไม่งั้นอ่านได้ข้อความของรอบก่อนที่ยังค้างบนจอ 10 วินาที (J1 แดงหลอก 29/10)"""
    p.evaluate("() => { const t = document.querySelector('#toast'); if (t) { t.classList.remove('show'); t.textContent = ''; } }")
    p.click("#btn-save")
    return toast(p)


def open_row(p, pro):
    p.click("#tab-records"); p.fill("#rec-search", pro); p.wait_for_timeout(400)
    p.click("#rec-table tbody tr[data-id]"); p.click("#btn-edit")


def hist_n(rid):
    return dv("GET", f"le_histories?$select=le_historyid&$count=true&$filter=_le_record_value eq {rid}")["@odata.count"]


def new_bug_row(pro):
    body = {"le_logbookdate": "2026-10-01", "le_pronumber": pro, "le_sitecode": "BKK4874A", "le_sitecodenocand": "BKK4874", "le_name": "J",
            "le_duedate": "2026-10-15", "le_amount": 10, "le_paymentmethod": "CHEQUE", "le_accountant": "a", "le_requestor": "r",
            "le_team": "IBC", "le_system": "DASHBOARD", "le_source": "Manual"}
    req = urllib.request.Request(f"{BASE}/api/data/v9.2/le_logelectrics", method="POST", data=json.dumps(body).encode(),
                                 headers={"Authorization": f"Bearer {os.environ['PVDESK_TOKEN']}", "Content-Type": "application/json", "Prefer": "return=representation"})
    with urllib.request.urlopen(req) as r:
        return json.loads(r.read())["le_logelectricid"]


def csv_rows(rows, enc):
    out = io.StringIO(); wr = csv.writer(out); wr.writerow(HEADS)
    for r in rows:
        wr.writerow([v.strftime("%d/%m/%Y") if isinstance(v, datetime.datetime) else ("" if v is None else v) for v in (r[h] for h in HEADS)])
    return out.getvalue().encode(enc)


def group_h(b, srv):
    """ไล่บัครอบ 3 เฟส 01 ไฟล์จากเครื่องลูกทีม"""
    import subprocess, time as _t
    ctx, p = open_page(b, srv)
    # H1 CSV ที่เซฟจาก Excel ภาษาไทย (ANSI รหัส 874) ชื่อไทยต้องไม่เพี้ยน
    name = "ผู้ให้เช่า ทดสอบ ภาษาไทย"
    (TMP / "h1.csv").write_bytes(csv_rows([row(0, **{"PRO NUMBER": "BUG-H1", "NAME": name})], "cp874"))
    t = upload(p, TMP / "h1.csv")
    if p.is_visible("#btn-confirm-import"):
        p.click("#btn-confirm-import"); p.click("dialog.dlg[open] .btn.dark"); p.wait_for_selector("#imp-done, #imp-failed", timeout=60000)
    got = (dv("GET", "le_logelectrics?$select=le_name&$filter=le_pronumber eq 'BUG-H1'")["value"] or [{}])[0].get("le_name")
    check("H1 CSV ภาษาไทยแบบ ANSI (874) ชื่อไทยเข้า Dataverse ไม่เพี้ยน", got == name, f"{got!r} | {t[:80]}")
    # H2 ไฟล์ที่จัดรูปแบบลากยาวถึงแถวสุดท้ายของ Excel (1,048,576) มีข้อมูลจริง 2 แถว
    wb = Workbook(); ws = wb.active; ws.title = "LOG ELECTRIC"; ws.append(HEADS)
    for r in (row(1), row(2)): ws.append([r[h] for h in HEADS])
    from openpyxl.styles import PatternFill
    for c in range(1, 14): ws.cell(1048576, c).fill = PatternFill("solid", fgColor="FFFF00")
    wb.save(TMP / "h2.xlsx")
    t0 = _t.time(); t = upload(p, TMP / "h2.xlsx"); dt = _t.time() - t0
    check("H2 ไฟล์จัดรูปแบบยาวถึงแถวล้าน อ่านเร็ว (ไม่เกิน 10 วินาที) และนับแค่ 2 แถวจริง", "All 2 rows are valid" in t and dt < 10, f"{dt:.1f} วินาที | {t[:80]}")
    # H3 CSV ที่ Excel แปลงเลข PRO เป็น E+
    (TMP / "h3.csv").write_bytes(csv_rows([row(3, **{"PRO NUMBER": "5.10001E+14"})], "utf-8-sig"))
    with p.expect_download():
        t = upload(p, TMP / "h3.csv")
    check("H3 CSV ที่เลข PRO กลายเป็น 5.10001E+14 ถูกตีกลับพร้อมบอกวิธีแก้", "rejected" in t and "scientific" in t, t[:200])
    # H4 ไฟล์ .xls รุ่นเก่า (Excel 97-2003)
    src = TMP / "h4.xlsx"; xlsx("h4.xlsx", [row(4), row(5)])
    js = ("const vm=require('vm'),fs=require('fs');const c={};vm.createContext(c);vm.runInContext(fs.readFileSync(process.argv[1],'utf8'),c);"
          "const X=c.XLSX;const wb=X.read(fs.readFileSync(process.argv[2]),{type:'buffer',cellDates:true});fs.writeFileSync(process.argv[3],Buffer.from(X.write(wb,{type:'array',bookType:'biff8'})));")
    subprocess.run(["node", "-e", js, str(pathlib.Path(__file__).parents[1] / "web/vendor/xlsx.full.min.js"), str(src), str(TMP / "h4.xls")], check=True)
    t = upload(p, TMP / "h4.xls")
    check("H4 ไฟล์ .xls รุ่นเก่าอ่านได้ 2 แถว", "All 2 rows are valid" in t, t[:150])
    p.wait_for_load_state("networkidle")
    check("H ไม่มี error ในคอนโซล", not p.errors, "; ".join(p.errors[:3]))
    ctx.close()


def group_j(b, srv):
    """ไล่บัครอบ 3 เฟส 03 คนใช้พร้อมกันและบัญชี (พี่ปอนด์ 29/10 "เจาะอีก ๆ")"""
    import threading
    ctx, p = open_page(b, srv)
    # J1 สองคนแก้แถวเดียวกัน: เราเปิดฟอร์มค้างไว้ อีกคนบันทึกก่อน เรากดบันทึกทีหลัง
    rid = dv("GET", "le_logelectrics?$select=le_logelectricid&$filter=le_pronumber eq 'PRO100010'")["value"][0]["le_logelectricid"]
    open_row(p, "PRO100010"); p.fill("#f-remark", "J1 from tab A")
    dv("PATCH", f"le_logelectrics({rid})", {"le_remark": "J1 other person"})
    h0 = hist_n(rid)
    t = save(p)
    now = dv("GET", f"le_logelectrics({rid})?$select=le_remark")["le_remark"]
    check("J1 คนอื่นบันทึกก่อน เรากดบันทึกทีหลัง ขึ้นข้อความให้รีเฟรช ไม่ทับของคนอื่น ไม่มีประวัติเพิ่ม", "Someone else changed this row" in t and now == "J1 other person" and hist_n(rid) == h0,
          f"{t} | ค่าใน Dataverse {now} | ประวัติ {h0} -> {hist_n(rid)}")
    p.click("#btn-refresh"); p.wait_for_timeout(2500)
    open_row(p, "PRO100010"); p.fill("#f-remark", "J1 after refresh"); t = save(p)
    check("J1 กดรีเฟรชแล้วแก้ใหม่ บันทึกได้", t.startswith("Saved"), t)
    dv("PATCH", f"le_logelectrics({rid})", {"le_remark": None})
    # J2 แถวที่เปิดแก้อยู่ถูกลบไปแล้ว
    jid = new_bug_row("BUG-J2")
    p.click("#btn-refresh"); p.wait_for_timeout(2500)
    open_row(p, "BUG-J2"); p.fill("#f-remark", "x")
    dv("DELETE", f"le_logelectrics({jid})")
    t = save(p)
    check("J2 แถวที่เปิดแก้ถูกลบไปแล้ว ขึ้นข้อความอ่านเข้าใจ ว่าแถวไม่มีแล้ว ไม่ใช่ข้อความดิบของระบบ", "no longer exists" in t and "Does Not Exist" not in t, t)
    p.keyboard.press("Escape")
    # J3 session หมดอายุ , J4 เน็ตหลุด ตอนบันทึก
    open_row(p, "PRO100011"); p.fill("#f-remark", "x")
    p.route("**/api/data/v9.2/$batch", lambda route: route.fulfill(status=401, content_type="application/json", body=json.dumps({"error": {"code": "0x80072560", "message": "The user is not a member of the organization."}})))
    t = save(p)
    check("J3 session หมดอายุตอนบันทึก บอกให้รีเฟรชแล้วล็อกอินใหม่", "session" in t.lower() and "sign in" in t.lower(), t)
    p.unroute("**/api/data/v9.2/$batch")
    p.route("**/api/data/v9.2/$batch", lambda route: route.abort("internetdisconnected"))
    t = save(p)
    check("J4 เน็ตหลุดตอนบันทึก บอกว่าต่อเซิร์ฟเวอร์ไม่ได้ ไม่ใช่ Failed to fetch", "connection" in t.lower() and "failed to fetch" not in t.lower(), t)
    p.unroute("**/api/data/v9.2/$batch")
    p.keyboard.press("Escape"); p.wait_for_load_state("networkidle")
    # J4b เน็ตหลุดกลางการนำเข้า 250 แถว (ก้อนละ 100) ก้อนที่ 3 หลุด
    seen = {"n": 0}
    def flaky(route):
        seen["n"] += 1
        return route.abort("internetdisconnected") if seen["n"] == 3 else route.fallback()
    upload(p, xlsx("j4b.xlsx", [row(i, **{"PRO NUMBER": f"BUG-J4-{i}"}) for i in range(250)]))
    p.route("**/api/data/v9.2/$batch", flaky)
    p.click("#btn-confirm-import"); p.click("dialog.dlg[open] .btn.dark")
    p.wait_for_selector("#imp-failed, #imp-done", timeout=120000)
    t = p.inner_text("#imp-result"); p.unroute("**/api/data/v9.2/$batch")
    n = count(flt="startswith(le_pronumber,'BUG-J4-')")
    check("J4b เน็ตหลุดกลางการนำเข้า บอกจำนวนที่บันทึกไปแล้วตรงกับใน Dataverse และบอกวิธีหาแถวจาก SOURCE", "200 rows saved" in t and n == 200 and "j4b.xlsx" in t, f"{t[:200]} | Dataverse {n}")
    p.wait_for_load_state("networkidle")
    ctx.close()
    # J5 คนที่ไม่มีบทบาท LE เลย
    u64 = user_id("64050580")
    def no_role(pg):
        pg.route("**/systemuserroles_association*", lambda r: r.fulfill(status=200, content_type="application/json", body=json.dumps({"value": [{"name": "PV Member"}]})))
        pg.route("**/teammembership_association*", lambda r: r.fulfill(status=200, content_type="application/json", body=json.dumps({"value": []})))
    ctx, p = open_page(b, srv, caller=u64, pre=no_role)
    t = p.inner_text("body")
    check("J5 คนที่ไม่มีบทบาท LE เลย เห็นข้อความว่ายังไม่มีสิทธิ์ ไม่เห็นตาราง", "access" in t.lower() and not p.query_selector("#rec-table"), t[:200])
    p.screenshot(path=str(EVID / "bug_j5_no_role.png")); ctx.close()
    # J6 ได้บทบาทผ่าน Team
    def team_mgr(pg):
        pg.route("**/systemuserroles_association*", lambda r: r.fulfill(status=200, content_type="application/json", body=json.dumps({"value": [{"name": "Basic User"}]})))
        pg.route("**/teammembership_association*", lambda r: r.fulfill(status=200, content_type="application/json", body=json.dumps({"value": [{"teamid": "x", "teamroles_association": [{"name": "LE Manager"}]}]})))
    ctx, p = open_page(b, srv, pre=team_mgr)
    check("J6 ได้บทบาท LE Manager ผ่าน Team แอปเปิดเป็น Manager", p.query_selector("#my-role") is not None and p.inner_text("#my-role") == "Manager", p.inner_text("body")[:120])
    ctx.close()
    # J7 สองคนเพิ่มแถวพร้อมกัน INDEX ต้องไม่ซ้ำ
    errs = []
    def burst(caller, tag):
        for i in range(40):
            try:
                dv("POST", "le_logelectrics", {"le_logbookdate": "2026-10-01", "le_pronumber": f"BUG-J7-{tag}{i}", "le_sitecode": "BKK4874A", "le_name": "J7",
                                               "le_duedate": "2026-10-15", "le_amount": 1, "le_paymentmethod": "CHEQUE", "le_accountant": "a", "le_requestor": "r",
                                               "le_team": "IBC", "le_system": "DASHBOARD", "le_source": "J7"}, caller=caller)
            except Exception as e:  # noqa: BLE001
                errs.append(str(e)[:80])
    th = [threading.Thread(target=burst, args=(None, "A")), threading.Thread(target=burst, args=(u64, "B"))]
    for x in th: x.start()
    for x in th: x.join()
    idx = [r["le_index"] for r in dv("GET", "le_logelectrics?$select=le_index&$filter=" + urllib.parse.quote("startswith(le_pronumber,'BUG-J7-')"))["value"]]
    check("J7 สองคนเพิ่มแถวพร้อมกัน 80 แถว INDEX ไม่ซ้ำกันเลย", len(idx) == 80 and len(set(idx)) == 80 and not errs, f"{len(idx)} แถว ไม่ซ้ำ {len(set(idx))} error {errs[:2]}")


def group_k(b, srv):
    """ไล่บัครอบ 3 เฟส 02 กับ 04: ชนเพดาน Dataverse (429) , จอโน้ตบุ๊กซูม 125% , ข้อความเป็นโค้ด HTML"""
    ctx, p = open_page(b, srv)
    # K1 นำเข้า 300 แถว Dataverse ตอบ 429 ก้อนที่ 2 กับ 3 (รอ 1 วินาที) แอปต้องรอแล้วไปต่อจนครบ ไม่หาย ไม่ซ้ำ
    seen = {"n": 0}
    def busy(route):
        seen["n"] += 1
        if seen["n"] in (2, 3):
            return route.fulfill(status=429, headers={"Retry-After": "1"}, content_type="application/json", body=json.dumps({"error": {"code": "0x80072322", "message": "Number of requests exceeded the limit"}}))
        return route.fallback()
    upload(p, xlsx("k1.xlsx", [row(i, **{"PRO NUMBER": f"BUG-K1-{i}"}) for i in range(300)]))
    p.route("**/api/data/v9.2/$batch", busy)
    p.click("#btn-confirm-import"); p.click("dialog.dlg[open] .btn.dark")
    p.wait_for_selector("#imp-done, #imp-failed", timeout=180000)
    p.unroute("**/api/data/v9.2/$batch")
    pros = [r["le_pronumber"] for r in dv("GET", "le_logelectrics?$select=le_pronumber&$filter=" + urllib.parse.quote("startswith(le_pronumber,'BUG-K1-')"))["value"]]
    check("K1 Dataverse ตอบ 429 สองครั้งระหว่างนำเข้า 300 แถว แอปรอแล้วไปต่อจนครบ ไม่หาย ไม่ซ้ำ", p.is_visible("#imp-done") and len(pros) == 300 and len(set(pros)) == 300 and seen["n"] >= 5,
          f"{len(pros)} แถว ไม่ซ้ำ {len(set(pros))} คำขอ batch {seen['n']} | {p.inner_text('#imp-result')[:80]}")
    p.wait_for_load_state("networkidle")
    # K4 ไม่มีปุ่ม Show more เลื่อนตารางลงสุดแล้วโหลดเพิ่มเอง (พี่ปอนด์ 29/10 ภาพ #75)
    p.click("#tab-records"); p.fill("#rec-search", ""); p.wait_for_timeout(400)
    n0 = p.eval_on_selector_all("#rec-table tbody tr[data-id]", "els => els.length")
    total = int(p.inner_text("#rec-summary").split(" of ")[0].replace(",", ""))
    for _ in range(6):
        p.evaluate("() => { const w = document.querySelector('#rec-table'); w.scrollTop = w.scrollHeight; window.scrollTo(0, document.body.scrollHeight); w.dispatchEvent(new Event('scroll')); window.dispatchEvent(new Event('scroll')); }")
        p.wait_for_timeout(400)
    n1 = p.eval_on_selector_all("#rec-table tbody tr[data-id]", "els => els.length")
    check("K4 ไม่มีปุ่ม Show more และเลื่อนลงสุดแล้วตารางโหลดแถวเพิ่มเองจนครบ", not p.query_selector("#rec-more") and total > 200 and n0 <= 200 and n1 == total, f"ก่อน {n0} หลัง {n1} ทั้งหมด {total}")
    # K3 ชื่อเป็นโค้ด HTML ต้องไม่ทำงาน
    evil = '<img src=x onerror="window.__pwn=1">'
    upload(p, xlsx("k3.xlsx", [row(0, **{"PRO NUMBER": "BUG-K3", "NAME": evil, "REMARK": "<b>bold</b>"})]))
    p.click("#btn-confirm-import"); p.click("dialog.dlg[open] .btn.dark"); p.wait_for_selector("#imp-done", timeout=60000)
    p.click("#tab-records"); p.fill("#rec-search", "BUG-K3"); p.wait_for_timeout(500)
    p.click("#rec-table tbody tr[data-id]"); p.wait_for_timeout(500)
    p.click('#rec-table th[data-col="name"] .xf-btn', force=True); p.wait_for_timeout(300); p.keyboard.press("Escape")
    if p.is_visible(".side-panel.open .sp-close"): p.click(".side-panel.open .sp-close")
    p.wait_for_timeout(300)
    p.select_option("#rec-group", "team"); p.wait_for_timeout(300)  # Group by ไม่มีตัวเลือก NAME ใช้ TEAM ให้หัวกลุ่มกับแถวโค้ด HTML ขึ้นพร้อมกัน
    p.select_option("#rec-group", "")
    pwn = p.evaluate("window.__pwn === 1")
    shown = evil in p.inner_text("#rec-table")
    check("K3 ชื่อเป็นโค้ด HTML แสดงเป็นตัวอักษร ไม่ทำงานทั้งในตาราง แผง ตัวกรอง และ Group by", not pwn and shown and not p.query_selector("#rec-table img, .side-panel img, .xf-pop img"), f"ทำงาน {pwn} แสดงเป็นตัวอักษร {shown}")
    p.wait_for_load_state("networkidle"); ctx.close()
    # K2 จอโน้ตบุ๊ก 1366x768 ซูม 125% ในกรอบ MDA ได้พื้นที่ราว 1093x560
    ctx, p = open_page(b, srv, vp=(1093, 560))
    m = p.evaluate("""() => { const bar = document.querySelector('.rec-tools, .toolbar, #records .row') || document.querySelector('#btn-add').parentElement;
      return { over: document.documentElement.scrollWidth > innerWidth + 1, addVisible: (() => { const r = document.querySelector('#btn-add').getBoundingClientRect(); return r.right <= innerWidth && r.bottom <= innerHeight; })(),
        rows: document.querySelectorAll('#rec-table tbody tr[data-id]').length, tableTop: Math.round(document.querySelector('#rec-table').getBoundingClientRect().top) }; }""")
    p.screenshot(path=str(EVID / "bug_k2_laptop_125.png"))
    check("K2 จอโน้ตบุ๊กซูม 125% ในกรอบ MDA (1093x560) ไม่ล้นแนวนอน ปุ่ม Add row อยู่ในจอ ตารางเริ่มไม่ต่ำเกินครึ่งจอ", not m["over"] and m["addVisible"] and m["tableTop"] < 280, str(m))
    ctx.close()


def cleanup():
    for r in dv("GET", "le_logelectrics?$select=le_logelectricid&$filter=" + urllib.parse.quote("startswith(le_pronumber,'BUG-')"))["value"]:
        dv("DELETE", f"le_logelectrics({r['le_logelectricid']})")
    r = dv("GET", "le_logelectrics?$select=le_logelectricid&$filter=le_pronumber eq 'PRO100005'")["value"]
    if r:
        dv("PATCH", f"le_logelectrics({r[0]['le_logelectricid']})", {"le_remark": None})


def main():
    srv = Server()
    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch()
            for g, fn in (("A", group_a), ("B", group_b), ("C", group_c), ("D", group_d), ("E", group_e), ("F", group_f), ("G", group_g), ("H", group_h), ("J", group_j), ("K", group_k)):
                if g not in GROUPS:
                    continue
                try:
                    fn(b, srv)
                except Exception as e:  # เคสล้มกลางทางนับเป็นข้อตก แล้วไปกลุ่มต่อไป
                    msg = str(e); wait = [l.strip() for l in msg.splitlines() if "waiting for" in l][:1]  # บอกว่ารอตัวไหน ไม่พ่น header ที่มี token
                    check(f"{g} รันจบไม่สะดุด", False, f"{type(e).__name__}: {msg.splitlines()[0][:150]} {wait}")
            b.close()
    finally:
        srv.proc.terminate()
        cleanup()
    bad = [n for n, ok in RESULTS if not ok]
    print(f"\nรวม {len(RESULTS)} ข้อ ผ่าน {len(RESULTS) - len(bad)} ตก {len(bad)}")
    sys.exit(1 if bad else 0)  # เดิมจบ 0 เสมอ สายคำสั่งที่ต่อด้วย && ไม่หยุดแม้มีข้อตก (29/10)


if __name__ == "__main__":
    main()
