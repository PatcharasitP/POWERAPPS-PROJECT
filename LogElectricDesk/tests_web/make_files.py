"""สร้างไฟล์ทดสอบจากเทมเพลตจริง (ข้อมูลสมมติทั้งหมด): good.xlsx ถูกทุกแถว, bad.xlsx มีแถวผิดที่รู้คำตอบล่วงหน้า"""
import random, datetime, sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "tests_web" / "files"
OUT.mkdir(parents=True, exist_ok=True)
random.seed(7)
SITES = ["BKK4874A", "CMI9012TI", "BRR7401", "SSK0104", "KKN0523B", "PKT0138", "NTB9167Z", "BKF8635"]
TEAMS, SYSTEMS, PAYS, COS = ["IBC", "MACRO", "ONLINE"], ["SITE RENTAL", "DASHBOARD"], ["CHEQUE", "BANK TRANSFER"], ["1019", "1068", "1069", None]


def rows(n, start=datetime.date(2026, 5, 1)):
    out = []
    for i in range(n):
        d = start + datetime.timedelta(days=random.randint(0, 175))
        out.append([d, f"PRO{100000 + i}", random.choice(SITES), f"ผู้ให้เช่าทดสอบ {i % 9 + 1}", d + datetime.timedelta(days=15),
                    round(random.uniform(500, 25000), 2), random.choice(PAYS), f"Accountant {i % 4 + 1}", f"Requestor {i % 6 + 1}",
                    random.choice(TEAMS), random.choice(SYSTEMS), random.choice(COS), "ข้อมูลทดสอบ" if i % 5 == 0 else None])
    return out


def save(name, data):
    wb = load_workbook(ROOT / "tools" / "out" / "LOG_ELECTRIC_TEMPLATE.xlsx")
    ws = wb["LOG ELECTRIC"]
    for r, row in enumerate(data, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(OUT / name)


save("good.xlsx", rows(40))
bad = rows(6)
bad[1][5] = "สองพัน"          # แถว 3 AMOUNT
bad[2][9] = "DTAC"            # แถว 4 TEAM
bad[3][0] = None              # แถว 5 LOG BOOK DATE ว่าง
bad[3][2] = "ไม่มีรหัส"       # แถว 5 SITE CODE
bad[4][11] = "9999"           # แถว 6 COMPANY CODE
save("bad.xlsx", bad)
# CSV ตามหัวเทมเพลต วันที่แบบ dd/mm/yyyy (ทีมเซฟจาก Excel ภาษาไทย)
import csv
heads = ["LOG BOOK DATE", "PRO NUMBER", "SITE CODE CANDIDATE", "NAME", "DUE DATE", "AMOUNT", "PAYMENT METHOD", "ACCOUNTANT", "REQUESTOR", "TEAM", "SYSTEM", "COMPANY CODE", "REMARK"]
with open(OUT / "good.csv", "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(heads)
    for r in rows(5):
        w.writerow([v.strftime("%d/%m/%Y") if isinstance(v, datetime.date) else ("" if v is None else v) for v in r])
print("ok", OUT)
